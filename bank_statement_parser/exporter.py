"""Export transactions to CSV or Excel spreadsheets."""

import csv
import os
from typing import List, Optional

from .models import Transaction
from .combiner import compute_summary


FIELDNAMES = [
    "Date",
    "Description",
    "Amount",
    "Type",
    "Balance",
    "Category",
    "Reference",
    "Bank",
    "Account",
    "Source File",
]


def export_csv(
    transactions: List[Transaction],
    output_path: str,
    include_summary: bool = False,
) -> str:
    """Export transactions to a CSV file.

    Args:
        transactions: List of transactions to export.
        output_path: Destination file path.
        include_summary: If True, append a summary section at the bottom.

    Returns:
        The absolute path of the written file.
    """
    output_path = os.path.abspath(output_path)

    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        for txn in transactions:
            writer.writerow(txn.to_dict())

        if include_summary:
            _write_csv_summary(writer, f, transactions)

    return output_path


def _write_csv_summary(writer, f, transactions: List[Transaction]) -> None:
    """Append summary rows to the CSV."""
    summary = compute_summary(transactions)
    f.write("\n")  # blank line separator
    f.write(f"# Summary\n")
    f.write(f"# Total Transactions: {summary['total_transactions']}\n")
    f.write(f"# Total Debits: {summary['total_debits']}\n")
    f.write(f"# Total Credits: {summary['total_credits']}\n")
    f.write(f"# Net: {summary['net']}\n")
    if summary["date_range"]:
        f.write(f"# Date Range: {summary['date_range'][0]} to {summary['date_range'][1]}\n")
    if summary["banks"]:
        f.write(f"# Banks: {', '.join(summary['banks'])}\n")


def export_excel(
    transactions: List[Transaction],
    output_path: str,
    include_summary: bool = True,
) -> str:
    """Export transactions to an Excel (.xlsx) file with formatting.

    Args:
        transactions: List of transactions to export.
        output_path: Destination file path.
        include_summary: If True, add a summary sheet.

    Returns:
        The absolute path of the written file.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, numbers, Alignment
    except ImportError:
        raise ImportError(
            "The 'openpyxl' package is required for Excel export. "
            "Install it with: pip install openpyxl"
        )

    output_path = os.path.abspath(output_path)
    wb = openpyxl.Workbook()

    # ── Transactions sheet ───────────────────────────────────────────────
    ws = wb.active
    ws.title = "Transactions"

    # Header row
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col_idx, name in enumerate(FIELDNAMES, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    debit_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    credit_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    for row_idx, txn in enumerate(transactions, start=2):
        data = txn.to_dict()
        row_fill = debit_fill if txn.is_debit else credit_fill
        for col_idx, field_name in enumerate(FIELDNAMES, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=data[field_name])
            cell.fill = row_fill
            # Format amount and balance as currency
            if field_name in ("Amount", "Balance") and isinstance(data[field_name], (int, float)):
                cell.number_format = '#,##0.00'

    # Auto-fit column widths
    for col_idx, name in enumerate(FIELDNAMES, start=1):
        max_len = len(name)
        for row_idx in range(2, min(len(transactions) + 2, 100)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 2, 40)

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(FIELDNAMES))}{len(transactions) + 1}"

    # ── Summary sheet ────────────────────────────────────────────────────
    if include_summary:
        _write_excel_summary(wb, transactions)

    wb.save(output_path)
    return output_path


def _write_excel_summary(wb, transactions: List[Transaction]) -> None:
    """Add a summary sheet to the workbook."""
    from openpyxl.styles import Font, PatternFill, Alignment

    summary = compute_summary(transactions)
    ws = wb.create_sheet("Summary")

    title_font = Font(bold=True, size=14)
    label_font = Font(bold=True)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    ws.cell(row=1, column=1, value="Bank Statement Summary").font = title_font

    rows = [
        ("Total Transactions", summary["total_transactions"]),
        ("Total Debits", summary["total_debits"]),
        ("Total Credits", summary["total_credits"]),
        ("Net Amount", summary["net"]),
    ]
    if summary["date_range"]:
        rows.append(("Date Range", f"{summary['date_range'][0]} to {summary['date_range'][1]}"))
    if summary["banks"]:
        rows.append(("Banks", ", ".join(summary["banks"])))
    if summary["source_files"]:
        rows.append(("Source Files", ", ".join(summary["source_files"])))

    for i, (label, value) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label).font = label_font
        cell = ws.cell(row=i, column=2, value=value)
        if isinstance(value, float):
            cell.number_format = '#,##0.00'

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 40
